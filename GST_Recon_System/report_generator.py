"""Excel reporting utilities for GST reconciliation outputs."""
from __future__ import annotations

from pathlib import Path
from typing import Dict

import pandas as pd
from openpyxl.styles import Font, PatternFill
from openpyxl.formatting.rule import CellIsRule


class ReportGenerator:
    """Writes multi-sheet reconciliation workbooks per client."""

    def __init__(self, logger):
        self.logger = logger

    def generate(self, client_name: str, result, output_dir: Path) -> Path:
        output_dir.mkdir(parents=True, exist_ok=True)
        output_path = output_dir / f"{client_name}_Reconciliation_Report.xlsx"

        sheets = {
            "SUMMARY": self._build_summary_sheet(result.summary),
            "MATCHED": result.matched,
            "MISMATCH": result.mismatched,
            "MISSING": result.missing,
            "EXTRA": result.extra,
        }

        with pd.ExcelWriter(output_path, engine="openpyxl") as writer:
            for name, data in sheets.items():
                data.to_excel(writer, sheet_name=name, index=False)
                worksheet = writer.sheets[name]
                self._style_header(worksheet)
                self._freeze_header(worksheet)
                self._auto_fit_columns(worksheet)

                if name == "MISMATCH":
                    self._highlight_mismatch(worksheet, data)

        self.logger.info("Report saved to %s", output_path)
        return output_path

    # ------------------------------------------------------------------
    # Sheet builders
    # ------------------------------------------------------------------

    def _build_summary_sheet(self, summary: Dict[str, float]) -> pd.DataFrame:
        summary_rows = [
            {"Metric": "Total Purchase Invoices", "Value": summary["total_purchase_invoices"]},
            {"Metric": "Total GSTR-2B Invoices", "Value": summary["total_gstr2b_invoices"]},
            {"Metric": "Matched Invoices", "Value": summary["matched_count"]},
            {"Metric": "Tax Mismatch Invoices", "Value": summary["mismatch_count"]},
            {"Metric": "Missing Invoices (Only in Purchase)", "Value": summary["missing_count"]},
            {"Metric": "Extra Invoices (Only in 2B)", "Value": summary["extra_count"]},
            {"Metric": "Total ITC as per Purchase", "Value": summary["total_itc_purchase"]},
            {"Metric": "Total ITC as per GSTR-2B", "Value": summary["total_itc_2b"]},
            {"Metric": "ITC Difference", "Value": summary["itc_difference"]},
        ]
        return pd.DataFrame(summary_rows)

    # ------------------------------------------------------------------
    # Styling helpers
    # ------------------------------------------------------------------

    def _style_header(self, worksheet):
        header_font = Font(bold=True)
        for cell in next(worksheet.iter_rows(min_row=1, max_row=1)):
            cell.font = header_font

    def _freeze_header(self, worksheet):
        worksheet.freeze_panes = "A2"

    def _auto_fit_columns(self, worksheet):
        for column_cells in worksheet.columns:
            max_length = 0
            column = column_cells[0].column_letter
            for cell in column_cells:
                if cell.value is None:
                    continue
                max_length = max(max_length, len(str(cell.value)))
            worksheet.column_dimensions[column].width = min(max_length + 2, 60)

    def _highlight_mismatch(self, worksheet, dataframe: pd.DataFrame):
        if dataframe.empty or "Tax_Difference" not in dataframe.columns:
            return

        try:
            diff_index = list(dataframe.columns).index("Tax_Difference") + 1
        except ValueError:
            return

        column_letter = worksheet.cell(row=1, column=diff_index).column_letter
        max_row = worksheet.max_row
        red_fill = PatternFill(start_color="FFC7CE", end_color="FFC7CE", fill_type="solid")

        worksheet.conditional_formatting.add(
            f"{column_letter}2:{column_letter}{max_row}",
            CellIsRule(operator="notEqual", formula=["0"], fill=red_fill),
        )
