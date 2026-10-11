"""Excel report generation for reconciliation results."""
from __future__ import annotations

from pathlib import Path
from typing import Dict

import pandas as pd
from openpyxl.styles import Font, PatternFill
from openpyxl.formatting.rule import CellIsRule

from app.utils.logger import logger


class ExcelReportBuilder:
    def __init__(self, output_dir: Path):
        self.output_dir = output_dir
        self.output_dir.mkdir(parents=True, exist_ok=True)

    def build(self, client_name: str, payload: Dict[str, pd.DataFrame], summary: Dict[str, float]) -> Path:
        output_path = self.output_dir / f"{client_name}_Reconciliation_Report.xlsx"

        frames = {
            "Summary": self._build_summary(summary),
            "Matched": payload["matched"],
            "Tax_Mismatch": payload["mismatch"],
            "Missing_in_2B": payload["missing"],
            "Extra_in_2B": payload["extra"],
        }

        with pd.ExcelWriter(output_path, engine="openpyxl") as writer:
            for sheet, df in frames.items():
                df.to_excel(writer, sheet_name=sheet, index=False)
                worksheet = writer.sheets[sheet]
                self._style_sheet(worksheet)
                if sheet == "Tax_Mismatch" and not df.empty:
                    self._highlight_mismatch(worksheet, df)

        logger.info("Report saved to {}", output_path)
        return output_path

    def _build_summary(self, summary: Dict[str, float]) -> pd.DataFrame:
        rows = [{"Metric": key.replace("_", " ").title(), "Value": value} for key, value in summary.items()]
        return pd.DataFrame(rows)

    def _style_sheet(self, worksheet):
        worksheet.freeze_panes = "A2"
        header_font = Font(bold=True)
        for cell in worksheet[1]:
            cell.font = header_font
        for column_cells in worksheet.columns:
            max_length = max((len(str(cell.value)) if cell.value else 0) for cell in column_cells)
            column_letter = column_cells[0].column_letter
            worksheet.column_dimensions[column_letter].width = min(max_length + 4, 60)

    def _highlight_mismatch(self, worksheet, df: pd.DataFrame):
        if "Tax_Difference" not in df.columns:
            return
        idx = list(df.columns).index("Tax_Difference") + 1
        column_letter = worksheet.cell(row=1, column=idx).column_letter
        red_fill = PatternFill(start_color="FFC7CE", end_color="FFC7CE", fill_type="solid")
        worksheet.conditional_formatting.add(
            f"{column_letter}2:{column_letter}{worksheet.max_row}",
            CellIsRule(operator="notEqual", formula=["0"], fill=red_fill),
        )
