"""Vectorized reconciliation logic.

Runs either inside a Celery worker (when Redis is configured) or in-process via
FastAPI BackgroundTasks (see ``dispatch_reconciliation``).
"""
from __future__ import annotations

from pathlib import Path
from typing import Dict, Iterable

import numpy as np
import pandas as pd
from openpyxl import load_workbook
from sqlalchemy.orm import Session

from app.celery_app import celery_app
from app.database import SessionLocal
from app.models import Client, JobStatus, ReconciliationJob
from app.services.report_generator import ExcelReportBuilder
from app.utils.file_validator import (
    REQUIRED_2B_COLUMNS,
    REQUIRED_PURCHASE_COLUMNS,
    ValidationError,
)
from app.utils.logger import logger
from app.utils.normalizer import normalize_dates, normalize_gstin, normalize_invoice

CHUNK_SIZE = 100_000
TOLERANCE = 1.0


def _iter_excel_chunks(file_path: Path, required_columns: Iterable[str]) -> Iterable[pd.DataFrame]:
    wb = load_workbook(file_path, read_only=True, data_only=True)
    try:
        ws = wb.active
        header_row = next(ws.iter_rows(min_row=1, max_row=1, values_only=True))
        header = [cell for cell in header_row]
        missing = [col for col in required_columns if col not in header]
        if missing:
            raise ValidationError(f"Missing required columns: {', '.join(missing)}")
        rows = []
        for row in ws.iter_rows(min_row=2, values_only=True):
            rows.append(row)
            if len(rows) >= CHUNK_SIZE:
                yield pd.DataFrame(rows, columns=header)
                rows = []
        if rows:
            yield pd.DataFrame(rows, columns=header)
    finally:
        wb.close()


def _prepare_purchase(df: pd.DataFrame) -> pd.DataFrame:
    optional_cols = ["CGST", "SGST", "IGST", "Taxable_Value"]
    for col in optional_cols:
        if col not in df.columns:
            df[col] = 0
    df["gstin"] = normalize_gstin(df["Supplier_GSTIN"])
    df["invoice_number"] = normalize_invoice(df["Invoice_No"])
    df["invoice_date"] = normalize_dates(df["Invoice_Date"])
    numeric_cols = ["Total_Tax", "CGST", "SGST", "IGST", "Taxable_Value"]
    for col in numeric_cols:
        if col not in df.columns:
            df[col] = 0
        df[col] = pd.to_numeric(df[col], errors="coerce").fillna(0.0)
    component_tax = df["CGST"] + df["SGST"] + df["IGST"]
    df["Total_Tax"] = df["Total_Tax"].where(df["Total_Tax"] != 0, component_tax)
    df = df[(df["gstin"] != "") & (df["invoice_number"] != "") & df["invoice_date"].notna()].copy()
    df["match_key"] = (
        df["gstin"].astype(str)
        + "|"
        + df["invoice_number"].astype(str)
        + "|"
        + df["invoice_date"].astype(str)
    )
    collapsed = df.groupby("match_key", as_index=False).agg(
        gstin=("gstin", "first"),
        invoice_number=("invoice_number", "first"),
        invoice_date=("invoice_date", "first"),
        purchase_taxable_value=("Taxable_Value", "sum"),
        purchase_total_tax=("Total_Tax", "sum"),
    )
    return collapsed


def _prepare_gstr2b(df: pd.DataFrame) -> pd.DataFrame:
    if "Taxable_Value" not in df.columns:
        df["Taxable_Value"] = 0
    if "Filing_Status" not in df.columns:
        df["Filing_Status"] = "UNKNOWN"
    df["gstin"] = normalize_gstin(df["GSTIN_of_Supplier"])
    df["invoice_number"] = normalize_invoice(df["Invoice_Number"])
    df["invoice_date"] = normalize_dates(df["Invoice_Date"])
    df["Taxable_Value"] = pd.to_numeric(df.get("Taxable_Value", 0), errors="coerce").fillna(0.0)
    df["Tax_Amount"] = pd.to_numeric(df["Tax_Amount"], errors="coerce").fillna(0.0)
    df["Filing_Status"] = df.get("Filing_Status", "UNKNOWN").fillna("UNKNOWN").astype(str).str.upper()
    df = df[(df["gstin"] != "") & (df["invoice_number"] != "") & df["invoice_date"].notna()].copy()
    df["match_key"] = (
        df["gstin"].astype(str)
        + "|"
        + df["invoice_number"].astype(str)
        + "|"
        + df["invoice_date"].astype(str)
    )
    collapsed = df.groupby("match_key", as_index=False).agg(
        gstin=("gstin", "first"),
        invoice_number=("invoice_number", "first"),
        invoice_date=("invoice_date", "first"),
        filing_status=("Filing_Status", "first"),
        gstr2b_taxable_value=("Taxable_Value", "sum"),
        gstr2b_tax_amount=("Tax_Amount", "sum"),
    )
    return collapsed


def _load_normalized(file_path: Path, required_columns, transform_func) -> pd.DataFrame:
    frames = []
    for chunk in _iter_excel_chunks(file_path, required_columns):
        frames.append(transform_func(chunk))
    if not frames:
        raise ValidationError("No records found after filtering")
    return pd.concat(frames, ignore_index=True)


def _merge_frames(purchase_df: pd.DataFrame, gstr2b_df: pd.DataFrame) -> Dict[str, pd.DataFrame]:
    merged = purchase_df.merge(
        gstr2b_df, on="match_key", how="outer", indicator=True, suffixes=("_purchase", "_gstr2b")
    )
    # Both sides carry gstin/invoice_number/invoice_date; keep whichever side is present.
    for field in ["gstin", "invoice_number", "invoice_date"]:
        merged[field] = merged[f"{field}_purchase"].combine_first(merged[f"{field}_gstr2b"])
        merged = merged.drop(columns=[f"{field}_purchase", f"{field}_gstr2b"])
    if "filing_status" not in merged.columns:
        merged["filing_status"] = "-"
    merged["filing_status"] = merged["filing_status"].fillna("-")
    for col in ["purchase_taxable_value", "purchase_total_tax", "gstr2b_taxable_value", "gstr2b_tax_amount"]:
        if col not in merged.columns:
            merged[col] = 0.0
        merged[col] = pd.to_numeric(merged[col], errors="coerce").fillna(0.0)
    merged["Tax_Difference"] = (
        merged["purchase_total_tax"].fillna(0) - merged["gstr2b_tax_amount"].fillna(0)
    )
    merged["status"] = np.select(
        [
            (merged["_merge"] == "both") & (merged["Tax_Difference"].abs() <= TOLERANCE),
            (merged["_merge"] == "both") & (merged["Tax_Difference"].abs() > TOLERANCE),
            merged["_merge"] == "left_only",
            merged["_merge"] == "right_only",
        ],
        ["MATCH", "TAX_MISMATCH", "MISSING_IN_2B", "EXTRA_IN_2B"],
        default="UNKNOWN",
    )
    views = {
        "matched": merged[merged["status"] == "MATCH"],
        "mismatch": merged[merged["status"] == "TAX_MISMATCH"],
        "missing": merged[merged["status"] == "MISSING_IN_2B"],
        "extra": merged[merged["status"] == "EXTRA_IN_2B"],
    }
    for key, frame in views.items():
        views[key] = frame[
            [
                "gstin",
                "invoice_number",
                "invoice_date",
                "purchase_taxable_value",
                "purchase_total_tax",
                "gstr2b_taxable_value",
                "gstr2b_tax_amount",
                "Tax_Difference",
                "filing_status",
            ]
        ].fillna("-")
    return views


def _summaries(views: Dict[str, pd.DataFrame], purchase_df: pd.DataFrame, gstr2b_df: pd.DataFrame) -> Dict[str, float]:
    totals = {
        "total_purchase_invoices": int(len(purchase_df)),
        "total_gstr2b_invoices": int(len(gstr2b_df)),
        "match_count": int(len(views["matched"])),
        "mismatch_count": int(len(views["mismatch"])),
        "missing_count": int(len(views["missing"])),
        "extra_count": int(len(views["extra"])),
        "total_itc_purchase": float(purchase_df["purchase_total_tax"].sum()),
        "total_itc_2b": float(gstr2b_df["gstr2b_tax_amount"].sum()),
    }
    totals["difference"] = totals["total_itc_purchase"] - totals["total_itc_2b"]
    return totals


def run_reconciliation(job_id: int) -> None:
    """Execute a reconciliation job synchronously and persist the results."""
    db: Session = SessionLocal()
    job = db.get(ReconciliationJob, job_id)
    if not job:
        logger.error("Job {} not found", job_id)
        db.close()
        return
    job.status = JobStatus.processing
    db.commit()

    client: Client = job.client
    storage_dir = Path(job.report_path or Path(job.purchase_path).parent)

    try:
        purchase_df = _load_normalized(Path(job.purchase_path), REQUIRED_PURCHASE_COLUMNS, _prepare_purchase)
        gstr2b_df = _load_normalized(Path(job.gstr2b_path), REQUIRED_2B_COLUMNS, _prepare_gstr2b)
        views = _merge_frames(purchase_df, gstr2b_df)
        summary = _summaries(views, purchase_df, gstr2b_df)

        builder = ExcelReportBuilder(storage_dir)
        report_path = builder.build(client.name.replace(" ", "_"), views, summary)

        job.status = JobStatus.completed
        job.total_invoices = int(summary["total_purchase_invoices"])
        job.match_count = summary["match_count"]
        job.mismatch_count = summary["mismatch_count"]
        job.missing_count = summary["missing_count"]
        job.extra_count = summary["extra_count"]
        job.purchase_itc_total = int(summary["total_itc_purchase"])
        job.gstr2b_itc_total = int(summary["total_itc_2b"])
        job.report_path = str(report_path)
        db.commit()
    except ValidationError as exc:
        logger.exception("Validation error for job {}", job_id)
        job.status = JobStatus.failed
        job.error_message = str(exc)
        db.commit()
    except Exception as exc:  # pragma: no cover
        logger.exception("Unexpected reconciliation failure for job {}", job_id)
        job.status = JobStatus.failed
        job.error_message = str(exc)
        db.commit()
    finally:
        db.close()


@celery_app.task(name="app.services.recon_engine.reconcile_job")
def reconcile_job(job_id: int) -> None:
    """Celery entry-point (used only when a Redis broker is configured)."""
    run_reconciliation(job_id)
